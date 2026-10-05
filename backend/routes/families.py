import csv
import io
import re
import uuid
from collections import defaultdict

from flask import Blueprint, request, jsonify, Response
from flask_jwt_extended import jwt_required
from openpyxl import Workbook
from openpyxl import load_workbook
from sqlalchemy import false, func, or_
from sqlalchemy.exc import IntegrityError
from flask_jwt_extended import get_jwt_identity

from models import (
    db,
    Barangay,
    Family,
    Resident,
    EvacuationAssignment,
    EvacuationCenterRoom,
    EvacuationCenter,
    SystemLog,
)

from utils import role_required


families_bp = Blueprint('families', __name__)

FAMILY_EXPORT_FIELDS = [
    'family_code', 'head_name', 'member_count', 'barangay',
    'contact_number', 'address', 'evacuation_center',
    'rooms', 'assigned_members', 'evacuation_status',
]


def _normalize_text(value):
    return re.sub(r'\s+', ' ', str(value or '').strip()).casefold()


def _family_query(search='', barangay_id='', assignment_status=''):
    query = Family.query.join(Barangay, Family.barangay_id == Barangay.id)
    search = str(search or '').strip()

    if search:
        pattern = f'%{search}%'
        query = query.filter(or_(
            Family.family_code.ilike(pattern),
            Family.head_name.ilike(pattern),
            Family.contact_number.ilike(pattern),
            Barangay.name.ilike(pattern),
            Family.id.in_(
                db.session.query(Resident.family_id).filter(
                    Resident.address.ilike(pattern),
                    Resident.family_id.is_not(None),
                )
            ),
        ))

    if barangay_id:
        try:
            query = query.filter(Family.barangay_id == int(barangay_id))
        except (TypeError, ValueError):
            query = query.filter(false())

    if assignment_status == 'assigned':
        query = query.filter(Family.id.in_(
            db.session.query(Resident.family_id)
            .join(EvacuationAssignment, EvacuationAssignment.resident_id == Resident.id)
            .filter(Resident.family_id.is_not(None))
        ))
    elif assignment_status == 'unassigned':
        query = query.filter(~Family.id.in_(
            db.session.query(Resident.family_id)
            .join(EvacuationAssignment, EvacuationAssignment.resident_id == Resident.id)
            .filter(Resident.family_id.is_not(None))
        ))

    return query


def _family_list_rows(families):
    family_ids = [family.id for family in families]
    if not family_ids:
        return []

    members = Resident.query.filter(Resident.family_id.in_(family_ids)).order_by(Resident.id.asc()).all()
    members_by_family = defaultdict(list)
    for member in members:
        members_by_family[member.family_id].append(member)

    assignment_rows = (
        db.session.query(
            Resident.family_id,
            EvacuationCenter.name,
            EvacuationCenterRoom.room_number,
            func.count(EvacuationAssignment.id),
        )
        .join(EvacuationAssignment, EvacuationAssignment.resident_id == Resident.id)
        .join(EvacuationCenterRoom, EvacuationAssignment.room_id == EvacuationCenterRoom.id)
        .join(EvacuationCenter, EvacuationCenterRoom.evacuation_center_id == EvacuationCenter.id)
        .filter(Resident.family_id.in_(family_ids))
        .group_by(Resident.family_id, EvacuationCenter.name, EvacuationCenterRoom.room_number)
        .all()
    )
    assignments_by_family = defaultdict(list)
    for family_id, center_name, room_number, assigned_count in assignment_rows:
        assignments_by_family[family_id].append({
            'center': center_name,
            'room': room_number,
            'assigned_members': assigned_count,
        })

    result = []
    for family in families:
        family_members = members_by_family.get(family.id, [])
        assignments = assignments_by_family.get(family.id, [])
        first_member = family_members[0] if family_members else None
        result.append({
            **family.to_dict(),
            'barangay': family.barangay.name if family.barangay else None,
            'member_count': family.member_count or len(family_members),
            'resident_count': len(family_members),
            'contact_number': family.contact_number or (first_member.contact_number if first_member else None),
            'address': next((member.address for member in family_members if member.address), None),
            'assignments': assignments,
            'evacuation_center': ', '.join(dict.fromkeys(item['center'] for item in assignments)) or None,
            'rooms': ', '.join(dict.fromkeys(item['room'] for item in assignments)) or None,
            'assigned_members': sum(item['assigned_members'] for item in assignments),
            'evacuation_status': 'Assigned' if assignments else family.evacuation_status or 'Not Assigned',
        })
    return result


@families_bp.route('/api/families', methods=['GET'])
@role_required('admin', 'lgu_personnel')
def list_families():
    search = request.args.get('search', '')
    barangay_id = request.args.get('barangay_id', '')
    assignment_status = request.args.get('assignment_status', '')
    sort_by = request.args.get('sort_by', 'family_code')
    direction = request.args.get('direction', 'asc')
    allowed_sorts = {
        'family_code': Family.family_code,
        'head_name': Family.head_name,
        'member_count': Family.member_count,
        'created_at': Family.created_at,
        'evacuation_status': Family.evacuation_status,
        'barangay': Barangay.name,
    }
    sort_column = allowed_sorts.get(sort_by, Family.family_code)
    order = sort_column.desc() if direction == 'desc' else sort_column.asc()
    query = _family_query(search, barangay_id, assignment_status)
    total = query.count()

    if request.args.get('all') == 'true':
        families = query.order_by(order, Family.id.asc()).all()
        return jsonify({'items': _family_list_rows(families), 'total': total})

    try:
        page = max(1, int(request.args.get('page', 1)))
        per_page = min(100, max(1, int(request.args.get('per_page', 20))))
    except (TypeError, ValueError):
        return jsonify({'error': 'page and per_page must be integers'}), 400

    families = query.order_by(order, Family.id.asc()).offset((page - 1) * per_page).limit(per_page).all()
    return jsonify({
        'items': _family_list_rows(families),
        'total': total,
        'page': page,
        'per_page': per_page,
        'pages': (total + per_page - 1) // per_page,
    })


@families_bp.route('/api/families/export', methods=['GET'])
@role_required('admin', 'lgu_personnel')
def export_families():
    query = _family_query(
        request.args.get('search', ''),
        request.args.get('barangay_id', ''),
        request.args.get('assignment_status', ''),
    )
    families = query.order_by(Family.family_code.asc(), Family.id.asc()).all()
    rows = _family_list_rows(families)
    output = io.StringIO(newline='')
    writer = csv.DictWriter(output, fieldnames=FAMILY_EXPORT_FIELDS)
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row.get(key) or '' for key in FAMILY_EXPORT_FIELDS})
    return Response(
        output.getvalue(),
        mimetype='text/csv; charset=utf-8',
        headers={'Content-Disposition': 'attachment; filename="evacway-families.csv"'},
    )


@families_bp.route('/api/families/template', methods=['GET'])
@role_required('admin', 'lgu_personnel')
def download_family_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Families'
    sheet.append([
        'Family Code', 'Head of Family', 'Barangay', 'Contact Number',
        'Address', 'Member First Name', 'Member Middle Name',
        'Member Last Name', 'Member Contact Number',
    ])
    instructions = workbook.create_sheet('Instructions')
    instructions.append(['One row per family member; repeat the family fields for each member.'])
    instructions.append(['Family Code must be unique and identical for all rows belonging to one family.'])
    instructions.append(['Required on every row: Family Code, Head of Family, Barangay, Member First Name, Member Last Name.'])
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        output.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="evacway-family-template.xlsx"'},
    )


def _split_full_name(full_name):
    parts = re.sub(r'\s+', ' ', str(full_name or '').strip()).split(' ')
    if not parts or not parts[0]:
        return None
    return parts[0], ' '.join(parts[1:-1]) or None, parts[-1]


def _family_duplicate(barangay_id, family_code, head_name):
    if Family.query.filter(func.lower(Family.family_code) == family_code.casefold()).first():
        return 'Family code already exists.'

    possible = Family.query.filter_by(barangay_id=barangay_id).all()
    if any(_normalize_text(item.head_name) == _normalize_text(head_name) for item in possible):
        return 'A family with this head already exists in the selected barangay.'
    return None


def _normalize_header(value):
    return re.sub(r'[^a-z0-9]+', '_', str(value or '').strip().casefold()).strip('_')


FAMILY_IMPORT_HEADERS = {
    'family_code': {'family_code', 'family_id', 'household_id'},
    'head_name': {'head_of_family', 'head_name', 'household_head_name'},
    'barangay': {'barangay', 'barangay_name'},
    'contact_number': {'contact_number', 'contact', 'family_contact'},
    'address': {'address', 'purok', 'address_purok'},
    'first_name': {'member_first_name', 'first_name'},
    'middle_name': {'member_middle_name', 'middle_name'},
    'last_name': {'member_last_name', 'last_name'},
    'member_contact': {'member_contact_number', 'member_contact'},
}


def _preview_family_workbook(file_storage):
    filename = str(file_storage.filename or '')
    if not filename.casefold().endswith('.xlsx'):
        return {'families': [], 'errors': [{'row': 0, 'message': 'Upload an .xlsx file.'}], 'rows': 0}

    try:
        workbook = load_workbook(io.BytesIO(file_storage.read()), read_only=True, data_only=True)
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        headers = next(values, None)
    except Exception:
        return {'families': [], 'errors': [{'row': 0, 'message': 'The workbook could not be read.'}], 'rows': 0}

    if not headers:
        return {'families': [], 'errors': [{'row': 1, 'message': 'The worksheet is empty.'}], 'rows': 0}

    column_indexes = {}
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        for field, aliases in FAMILY_IMPORT_HEADERS.items():
            if normalized in aliases:
                column_indexes[field] = index

    required = {'family_code', 'head_name', 'barangay', 'first_name', 'last_name'}
    missing = sorted(required - set(column_indexes))
    if missing:
        return {
            'families': [],
            'errors': [{'row': 1, 'message': f'Missing required columns: {", ".join(missing)}'}],
            'rows': 0,
        }

    errors = []
    families_by_code = {}
    row_count = 0

    def cell(row, key):
        index = column_indexes.get(key)
        if index is None or index >= len(row):
            return ''
        return re.sub(r'\s+', ' ', str(row[index] or '').strip())

    for row_number, row in enumerate(values, start=2):
        if not row or not any(value is not None and str(value).strip() for value in row):
            continue
        row_count += 1
        family_code = cell(row, 'family_code')
        head_name = cell(row, 'head_name')
        barangay_name = cell(row, 'barangay')
        first_name = cell(row, 'first_name')
        middle_name = cell(row, 'middle_name')
        last_name = cell(row, 'last_name')

        missing_fields = [
            label for value, label in (
                (family_code, 'Family Code'),
                (head_name, 'Head of Family'),
                (barangay_name, 'Barangay'),
                (first_name, 'Member First Name'),
                (last_name, 'Member Last Name'),
            ) if not value
        ]
        if missing_fields:
            errors.append({'row': row_number, 'message': f'Missing required values: {", ".join(missing_fields)}'})
            continue
        if len(family_code) > 100 or len(head_name) > 150:
            errors.append({'row': row_number, 'message': 'Family Code or Head of Family exceeds its maximum length.'})
            continue

        barangay = Barangay.query.filter(func.lower(Barangay.name) == barangay_name.casefold()).first()
        if not barangay:
            errors.append({'row': row_number, 'message': f'Barangay "{barangay_name}" is not registered.'})
            continue

        code_key = family_code.casefold()
        metadata = {
            'family_code': family_code,
            'head_name': head_name,
            'barangay_id': barangay.id,
            'barangay': barangay.name,
            'contact_number': cell(row, 'contact_number') or None,
            'address': cell(row, 'address') or None,
            'members': [],
            'source_rows': [],
        }
        existing_group = families_by_code.get(code_key)
        if existing_group is None:
            families_by_code[code_key] = metadata
            existing_group = metadata
        elif (
            existing_group['barangay_id'] != barangay.id
            or _normalize_text(existing_group['head_name']) != _normalize_text(head_name)
        ):
            errors.append({'row': row_number, 'message': 'Rows sharing a Family Code must have the same head and barangay.'})
            continue

        member_name = ' '.join(part for part in (first_name, middle_name, last_name) if part)
        member_key = _normalize_text(member_name)
        if any(_normalize_text(item['full_name']) == member_key for item in existing_group['members']):
            errors.append({'row': row_number, 'message': f'Duplicate member "{member_name}" for Family Code {family_code}.'})
            continue

        existing_group['members'].append({
            'first_name': first_name,
            'middle_name': middle_name or None,
            'last_name': last_name,
            'full_name': member_name,
            'contact_number': cell(row, 'member_contact') or existing_group['contact_number'],
        })
        existing_group['source_rows'].append(row_number)

    families = list(families_by_code.values())
    seen_identities = set()
    for family in families:
        identity = (family['barangay_id'], _normalize_text(family['head_name']))
        if identity in seen_identities:
            errors.append({'row': family['source_rows'][0], 'message': 'Duplicate head and barangay found with another Family Code in this workbook.'})
        seen_identities.add(identity)

        if not any(_normalize_text(member['full_name']) == _normalize_text(family['head_name']) for member in family['members']):
            errors.append({'row': family['source_rows'][0], 'message': 'The household head must also appear among this family’s member rows.'})

        duplicate = _family_duplicate(family['barangay_id'], family['family_code'], family['head_name'])
        if duplicate:
            errors.append({'row': family['source_rows'][0], 'message': duplicate})

        family['member_count'] = len(family['members'])

    return {'families': families, 'errors': errors, 'rows': row_count}


@families_bp.route('/api/families/import', methods=['POST'])
@role_required('admin', 'lgu_personnel')
def import_families():
    file_storage = request.files.get('file')
    if not file_storage:
        return jsonify({'error': 'Choose an .xlsx file to import.'}), 400

    preview = _preview_family_workbook(file_storage)
    commit_requested = str(request.form.get('commit', '')).casefold() == 'true'
    result = {
        'families': [
            {key: family[key] for key in ('family_code', 'head_name', 'barangay', 'contact_number', 'address', 'member_count')}
            for family in preview['families']
        ],
        'errors': preview['errors'],
        'rows': preview['rows'],
        'can_import': bool(preview['families']) and not preview['errors'],
    }

    if not commit_requested:
        return jsonify(result), 200
    if not result['can_import']:
        return jsonify({**result, 'error': 'Resolve all import errors before committing.'}), 400

    created_count = 0
    try:
        for item in preview['families']:
            family = Family(
                barangay_id=item['barangay_id'],
                family_code=item['family_code'],
                head_name=item['head_name'],
                member_count=item['member_count'],
                contact_number=item['contact_number'],
                evacuation_status='Not Assigned',
            )
            db.session.add(family)
            db.session.flush()
            for member in item['members']:
                db.session.add(Resident(
                    family_id=family.id,
                    first_name=member['first_name'],
                    middle_name=member['middle_name'],
                    last_name=member['last_name'],
                    household_head_name=item['head_name'],
                    barangay=item['barangay'],
                    address=item['address'],
                    contact_number=member['contact_number'],
                    family_members=item['member_count'],
                    uploaded_by=get_jwt_identity(),
                ))
            created_count += 1
        db.session.add(SystemLog(
            user_id=get_jwt_identity(),
            action=f'Imported {created_count} families from Excel',
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()
        return jsonify({'error': 'Import failed. No family records were committed.'}), 500

    return jsonify({
        'message': f'Imported {created_count} families ({preview["rows"]} member rows).',
        'successful_families': created_count,
        'successful_rows': preview['rows'],
        'failed_rows': 0,
    }), 201


@families_bp.route('/api/families/<int:family_id>', methods=['GET'])
@role_required('admin', 'lgu_personnel')
def get_family(family_id):
    family = Family.query.get_or_404(family_id)
    result = _family_list_rows([family])[0]
    result['members'] = [member.to_dict() for member in Resident.query.filter_by(family_id=family.id).order_by(Resident.id.asc()).all()]
    return jsonify(result)


@families_bp.route('/api/families', methods=['POST'])
@role_required('admin', 'lgu_personnel')
def add_family():
    data = request.get_json(silent=True) or {}
    try:
        barangay_id = int(data.get('barangay_id'))
    except (TypeError, ValueError):
        return jsonify({'error': 'Barangay is required'}), 400

    barangay = db.session.get(Barangay, barangay_id)
    if not barangay:
        return jsonify({'error': 'Barangay not found'}), 404

    head_name = re.sub(r'\s+', ' ', str(data.get('head_name') or '').strip())
    if not head_name or len(head_name) > 150:
        return jsonify({'error': 'Head of family is required and must be 150 characters or fewer'}), 400

    family_code = re.sub(r'\s+', ' ', str(data.get('family_code') or '').strip())
    if not family_code:
        family_code = f'FAM-{barangay_id}-{uuid.uuid4().hex[:8].upper()}'
    if len(family_code) > 100:
        return jsonify({'error': 'Family code must be 100 characters or fewer'}), 400

    contact_number = str(data.get('contact_number') or '').strip() or None
    address = str(data.get('address') or '').strip() or None
    if contact_number and len(contact_number) > 20:
        return jsonify({'error': 'Contact number must be 20 characters or fewer'}), 400
    if address and len(address) > 255:
        return jsonify({'error': 'Address must be 255 characters or fewer'}), 400

    extra_members = data.get('members') or []
    if not isinstance(extra_members, list) or len(extra_members) > 99:
        return jsonify({'error': 'Members must be a list of no more than 99 additional people'}), 400
    parsed_members = []
    member_keys = {_normalize_text(head_name)}
    for index, member in enumerate(extra_members, start=1):
        if not isinstance(member, dict):
            return jsonify({'error': f'Additional member {index} must include a full_name field'}), 400
        name = member.get('full_name')
        parsed_name = _split_full_name(name)
        if not parsed_name:
            return jsonify({'error': f'Additional member {index} needs a full name'}), 400
        if any(len(part) > 100 for part in parsed_name if part):
            return jsonify({'error': f'Additional member {index} has a name field longer than 100 characters'}), 400
        member_contact = str(member.get('contact_number') or '').strip()
        if len(member_contact) > 20:
            return jsonify({'error': f'Additional member {index} contact number must be 20 characters or fewer'}), 400
        if _normalize_text(name) in member_keys:
            return jsonify({'error': f'Duplicate member name: {name}'}), 400
        member_keys.add(_normalize_text(name))
        parsed_members.append((name.strip(), parsed_name))

    duplicate_error = _family_duplicate(barangay_id, family_code, head_name)
    if duplicate_error:
        return jsonify({'error': duplicate_error}), 409

    member_count = len(parsed_members) + 1
    family = Family(
        barangay_id=barangay_id,
        family_code=family_code,
        head_name=head_name,
        member_count=member_count,
        contact_number=contact_number,
        evacuation_status='Not Assigned',
    )

    try:
        db.session.add(family)
        db.session.flush()
        residents = [(head_name, _split_full_name(head_name), contact_number)]
        residents.extend(
            (name, parsed_name, str(member.get('contact_number') or '').strip() or contact_number)
            for member, (name, parsed_name) in zip(extra_members, parsed_members)
        )
        for full_name, name_parts, member_contact in residents:
            first_name, middle_name, last_name = name_parts
            db.session.add(Resident(
                family_id=family.id,
                first_name=first_name,
                middle_name=middle_name,
                last_name=last_name,
                household_head_name=head_name,
                barangay=barangay.name,
                address=address,
                contact_number=member_contact,
                family_members=member_count,
                uploaded_by=get_jwt_identity(),
            ))
        db.session.add(SystemLog(
            user_id=get_jwt_identity(),
            action=f'Added family {family.family_code} in {barangay.name}',
        ))
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({'error': 'A duplicate family record was detected; review the list and try again.'}), 409

    return jsonify({
        'message': 'Family added successfully',
        'family': _family_list_rows([family])[0],
    }), 201


# =========================================================
# ADD FAMILY MEMBER
# =========================================================
@families_bp.route(
    '/api/families/<int:family_id>/members',
    methods=['POST']
)
@jwt_required()
@role_required('admin', 'lgu_personnel')
def add_family_member(family_id):

    family = Family.query.get(family_id)

    if not family:
        return jsonify({
            "error": "Family not found"
        }), 404

    data = request.get_json()

    if not data:
        return jsonify({
            "error": "No data provided"
        }), 400

    household_head_name = data.get("household_head_name")
    barangay = data.get("barangay")
    address = data.get("address")
    contact_number = data.get("contact_number")
    vulnerable_members = data.get("vulnerable_members")
    remarks = data.get("remarks")

    if not household_head_name:
        return jsonify({
            "error": "Member name is required"
        }), 400

    if not barangay:
        barangay = family.barangay.name

    resident = Resident(
        family_id=family.id,
        household_head_name=household_head_name,
        barangay=barangay,
        address=address,
        contact_number=contact_number,
        family_members=1,
        vulnerable_members=vulnerable_members,
        remarks=remarks
    )

    db.session.add(resident)

    # Keep the family's member count updated
    family.member_count = Resident.query.filter_by(
        family_id=family.id
    ).count() + 1

    db.session.commit()

    return jsonify({
        "message": "Family member added successfully",
        "member": resident.to_dict()
    }), 201


# =========================================================
# GET FAMILY MEMBERS
# =========================================================
@families_bp.route(
    '/api/families/<int:family_id>/members',
    methods=['GET']
)
@jwt_required()
@role_required('admin', 'lgu_personnel')
def get_family_members(family_id):

    family = Family.query.get(family_id)

    if not family:
        return jsonify({
            "error": "Family not found"
        }), 404

    members = Resident.query.filter_by(
        family_id=family_id
    ).order_by(
        Resident.id.asc()
    ).all()

    return jsonify([
        member.to_dict()
        for member in members
    ]), 200